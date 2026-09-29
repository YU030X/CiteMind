"""自制 DOCX 样本构造器（无敏感内容、可离线复现、字节确定）。

只用 dev/worker 组已有的 ``python-docx`` 在内存中构造样本，不联网、不读取环境文件、不使用
任何真实企业资料。``_finalize`` 固定核心属性时间并重写 ZIP（条目排序、固定时间戳），使同一
Python/库版本下重复运行得到相同字节，便于单测与集成用例复用。

正样本 5 份（每份都只含收窄子集内的结构）：

1. ``multi_heading_interleaved``：多级标题 + 段落与简单表格交错；
2. ``simple_table``：普通表格，无合并；
3. ``horizontal_merge_grid_before``：横向合并 + ``gridBefore`` 偏移；
4. ``vertical_merge``：纵向合并，续格不复制上一行文字；
5. ``long_content``：长段落与长表格，用于验证超预算块的字符级切分。

负样本：嵌套表格、空文档、只有图片、宏部件、实体声明与损坏 ZIP；它们**不计入** 5 份可支持
正样本。
"""

from __future__ import annotations

import io
import struct
import zipfile
import zlib
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_FIXED_TIME = (2020, 1, 1, 0, 0, 0)


def _png_1x1() -> bytes:
    """生成一个合法的 1x1 RGB PNG，不依赖 PIL。"""

    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + tag
            + payload
            + struct.pack(">I", zlib.crc32(tag + payload) & 0xFFFFFFFF)
        )

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00\x00\x00\x00"  # filter=0 + 黑色像素
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _normalize_zip(data: bytes) -> bytes:
    """按名称排序、固定时间戳重写 ZIP，去掉 python-docx 的时钟与压缩差异。"""

    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=_FIXED_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return output.getvalue()


def _finalize(document: DocumentObject) -> bytes:
    document.core_properties.created = datetime(2020, 1, 1, tzinfo=UTC)
    document.core_properties.modified = datetime(2020, 1, 1, tzinfo=UTC)
    buffer = io.BytesIO()
    document.save(buffer)
    return _normalize_zip(buffer.getvalue())


def _set_grid_before(row: object, value: int) -> None:
    """给 ``w:tr`` 写 ``w:trPr/w:gridBefore``，模拟行首空网格列。"""

    tr_pr = row._tr.get_or_add_trPr()  # type: ignore[attr-defined]
    grid_before = OxmlElement("w:gridBefore")
    grid_before.set(qn("w:val"), str(value))
    tr_pr.append(grid_before)


def positive_multi_heading_interleaved() -> bytes:
    """多级标题 + 段落与简单表格交错，验证 heading_path 与段落索引。"""

    document = Document()
    document.add_heading("员工手册", level=1)
    document.add_paragraph("第一段正文。")
    document.add_heading("考勤", level=2)
    document.add_paragraph("第二段正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "上午"
    table.cell(0, 1).text = "09:00"
    table.cell(1, 0).text = "下午"
    table.cell(1, 1).text = "18:00"
    document.add_paragraph("表格之后的段落。")
    document.add_heading("补充", level=2)
    document.add_paragraph("最后一段。")
    return _finalize(document)


def positive_simple_table() -> bytes:
    """普通表格，无任何合并。"""

    document = Document()
    document.add_paragraph("报销标准：")
    table = document.add_table(rows=3, cols=3)
    header = ("项目", "上限", "备注")
    for column, value in enumerate(header):
        table.cell(0, column).text = value
    table.cell(1, 0).text = "市内交通"
    table.cell(1, 1).text = "50"
    table.cell(1, 2).text = "凭票"
    table.cell(2, 0).text = "住宿"
    table.cell(2, 1).text = "500"
    table.cell(2, 2).text = "按城市"
    return _finalize(document)


def positive_horizontal_merge_grid_before() -> bytes:
    """横向合并 + ``gridBefore`` 偏移，验证 origin 只出现一次与网格列偏移。"""

    document = Document()
    document.add_paragraph("合并示例：")
    table = document.add_table(rows=2, cols=3)
    table.cell(0, 0).text = "跨两列"
    table.cell(0, 2).text = "普通"
    table.cell(0, 0).merge(table.cell(0, 1))
    # 先把第二行文字写好，再去掉第一个 tc 并写 gridBefore=1：剩余两格落在网格列 2、3。
    table.cell(1, 0).text = "占位"
    table.cell(1, 1).text = "列二"
    table.cell(1, 2).text = "列三"
    second_row = table.rows[1]
    first_tc = second_row._tr.find(qn("w:tc"))
    second_row._tr.remove(first_tc)
    _set_grid_before(second_row, 1)
    return _finalize(document)


def positive_vertical_merge() -> bytes:
    """纵向合并：续格不复制上一行文字。"""

    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "跨行"
    table.cell(0, 1).text = "第一行"
    table.cell(1, 1).text = "第二行"
    table.cell(2, 1).text = "第三行"
    table.cell(0, 0).merge(table.cell(2, 0))
    return _finalize(document)


def positive_long_content() -> bytes:
    """长段落与长表格，用于验证超预算块的字符级切分与 block 内字符区间。"""

    document = Document()
    long_sentence = "这是一段用于验证切分的长文本，包含足够多的字符以便按预算拆分。" * 20
    document.add_paragraph(long_sentence)
    table = document.add_table(rows=12, cols=2)
    for row in range(12):
        table.cell(row, 0).text = f"字段{row}"
        table.cell(row, 1).text = f"取值{row}-" + "内容" * 6
    return _finalize(document)


def nested_table_docx() -> bytes:
    """单元格内含嵌套表格：刻意不支持，应静态失败而不是静默丢内容。"""

    document = Document()
    document.add_paragraph("外层：")
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "内层内容"
    return _finalize(document)


def empty_docx() -> bytes:
    """没有任何正文块的空文档。"""

    return _finalize(Document())


def image_only_docx() -> bytes:
    """只有图片、没有文字段落。"""

    document = Document()
    document.add_picture(io.BytesIO(_png_1x1()))
    return _finalize(document)


def macro_part_docx() -> bytes:
    """在普通 DOCX 上追加 ``word/vbaProject.bin`` 宏部件，应被拒绝。"""

    base = positive_simple_table()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as source:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.is_dir():
                    continue
                target.writestr(info, source.read(info.filename))
            target.writestr("word/vbaProject.bin", b"macro")
    return buffer.getvalue()


def entity_declaration_docx() -> bytes:
    """把 ``word/document.xml`` 前置一个实体声明，解析期应静态失败。"""

    base = positive_simple_table()
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as source:
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.is_dir():
                    continue
                data = source.read(info.filename)
                if info.filename == "word/document.xml":
                    data = data.replace(
                        b"<w:document",
                        b'<!DOCTYPE w:document [<!ENTITY x "boom">]><w:document',
                        1,
                    )
                target.writestr(info, data)
    return buffer.getvalue()


def corrupted_zip_docx() -> bytes:
    """合法 PK 魔数下的损坏 ZIP。"""

    return b"PK\x03\x04" + b"\x00" * 80


def positive_samples() -> dict[str, bytes]:
    """返回 5 份可支持正样本的稳定命名映射。"""

    return {
        "multi_heading_interleaved": positive_multi_heading_interleaved(),
        "simple_table": positive_simple_table(),
        "horizontal_merge_grid_before": positive_horizontal_merge_grid_before(),
        "vertical_merge": positive_vertical_merge(),
        "long_content": positive_long_content(),
    }


__all__ = [
    "corrupted_zip_docx",
    "empty_docx",
    "entity_declaration_docx",
    "image_only_docx",
    "macro_part_docx",
    "nested_table_docx",
    "positive_horizontal_merge_grid_before",
    "positive_long_content",
    "positive_multi_heading_interleaved",
    "positive_samples",
    "positive_simple_table",
    "positive_vertical_merge",
]
